"""Gratis deterministiese formuletoets: bou 'n werkboek met die item se toetsdata,
sit die formule(s) in, laat LibreOffice (headless) herbereken en vergelyk met die verwagte waarde.
Gebruik: python3 formula_check.py items.json  -> skryf formula_report.json, exit 1 by enige fout."""
import json,sys,os,subprocess,tempfile,csv,math
from openpyxl import Workbook
def run(item):
    t=item['toets']['formule']   # {"blaaie":{"Data":{"A1":5,...}}, "selle":[{"blad":"Data","sel":"B2","formule":"=...","verwag":12}]}
    wb=Workbook(); first=True
    for name,cells in t['blaaie'].items():
        ws=wb.active if first else wb.create_sheet(); ws.title=name; first=False
        for ref,v in cells.items(): ws[ref]=v
    for c in t['selle']: wb[c['blad']][c['sel']]=c['formule']
    d=tempfile.mkdtemp(); x=os.path.join(d,'t.xlsx'); wb.save(x)
    res=[]
    for c in t['selle']:
        idx=list(t['blaaie']).index(c['blad'])+1
        out=os.path.join(d,f"o{idx}")
        os.makedirs(out,exist_ok=True)
        subprocess.run(['soffice','--headless','--calc','--convert-to',f'csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,{idx}','--outdir',out,x],capture_output=True,timeout=120)
        f=[p for p in os.listdir(out) if p.endswith('.csv')][0]
        rows=list(csv.reader(open(os.path.join(out,f),encoding='utf-8')))
        from openpyxl.utils.cell import coordinate_from_string,column_index_from_string
        col,row=coordinate_from_string(c['sel']); got=rows[row-1][column_index_from_string(col)-1]
        exp=c['verwag']
        try: ok=math.isclose(float(got),float(exp),rel_tol=1e-9,abs_tol=1e-9)
        except ValueError: ok=(str(got)==str(exp))
        res.append({'sel':c['sel'],'formule':c['formule'],'verwag':exp,'kry':got,'ok':ok})
    return res
if __name__=='__main__':
    items=json.load(open(sys.argv[1],encoding='utf-8')); rep=[];fail=0
    for it in items:
        if 'formule' in it.get('toets',{}):
            r=run(it); ok=all(x['ok'] for x in r); fail+=not ok; rep.append({'id':it['id'],'ok':ok,'selle':r})
    json.dump(rep,open('formula_report.json','w'),ensure_ascii=False,indent=1)
    print(f"formule-items: {len(rep)}, foute: {fail}"); sys.exit(1 if fail else 0)
